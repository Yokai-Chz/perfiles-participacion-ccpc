"""Presentación ejecutiva (RE-03): 10 diapositivas, generada desde las tablas de `entrega/tablas/`."""
import polars as pl
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Cm, Pt

import contenido
from configuracion import NOMBRE_EQUIPO, TABLAS
from figuras import FIGURAS, SERIE

TINTA, TINTA_2, FONDO = RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0xFC, 0xFC, 0xFB)
ANCHO, ALTO = Cm(33.867), Cm(19.05)


def _rgb(hex_):
    return RGBColor.from_string(hex_.lstrip("#").upper())


def texto(diap, x, y, w, h, lineas, tam=18, color=TINTA, negrita=False):
    caja = diap.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    marco = caja.text_frame
    marco.word_wrap = True
    for i, linea in enumerate(lineas if isinstance(lineas, list) else [lineas]):
        par = marco.paragraphs[0] if i == 0 else marco.add_paragraph()
        par.text = linea
        par.space_after = Pt(6)
        for r in par.runs:
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Arial", Pt(tam), negrita, color
    return caja


def nueva(prs, titulo):
    diap = prs.slides.add_slide(prs.slide_layouts[6])
    diap.background.fill.solid()
    diap.background.fill.fore_color.rgb = FONDO
    texto(diap, 1.5, 0.9, 30.8, 2.4, titulo, tam=28, negrita=True)
    return diap


def imagen(diap, nombre, x, y, ancho):
    diap.shapes.add_picture(str(FIGURAS / nombre), Cm(x), Cm(y), width=Cm(ancho))


def tabla(diap, filas, x, y, anchos, tam=13):
    forma = diap.shapes.add_table(len(filas), len(filas[0]), Cm(x), Cm(y), Cm(sum(anchos)), Cm(1.2 * len(filas)))
    t = forma.table
    for j, a in enumerate(anchos):
        t.columns[j].width = Cm(a)
    for i, fila in enumerate(filas):
        for j, valor in enumerate(fila):
            c = t.cell(i, j)
            c.text = str(valor)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(0xEE, 0xED, 0xE9) if i == 0 else FONDO
            for par in c.text_frame.paragraphs:
                for r in par.runs:
                    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Arial", Pt(tam), i == 0, TINTA


def construir(ruta):
    m = contenido.metricas()
    cal = pl.read_csv(TABLAS / "calidad_nacional.csv")
    comp = pl.read_csv(TABLAS / "sensibilidad_comparacion.csv")
    ent, nac = comp.row(0, named=True), comp.row(1, named=True)
    filas_total = pl.read_csv(TABLAS / "tabla_calidad.csv")["filas"].sum()
    prs = Presentation()
    prs.slide_width, prs.slide_height = ANCHO, ALTO

    d = prs.slides.add_slide(prs.slide_layouts[6])
    d.background.fill.solid()
    d.background.fill.fore_color.rgb = FONDO
    texto(d, 2, 5.5, 29, 4, "Perfiles ciudadanos de participación", tam=40, negrita=True)
    texto(d, 2, 9.5, 29, 3, "Cinco maneras de participar en las elecciones federales, 2009 a 2024", tam=22, color=TINTA_2)
    texto(d, 2, 13.5, 29, 3, [f"Eje 1, Reto 1 · {NOMBRE_EQUIPO}", "Borrador para revisión del equipo"], tam=16, color=TINTA_2)

    d = nueva(prs, "Seis elecciones, 32 entidades y una pregunta: ¿qué tipos de ciudadanía existen?")
    texto(d, 1.5, 4, 14, 12, [f"Fuente: Conteos Censales de Participación Ciudadana (INE), 2009 a 2024.",
                              f"{filas_total:,.0f} filas agregadas por sección, sexo y edad; sin datos personales.",
                              f"Lista nominal de 2024: {cal.filter(pl.col('AELEC') == 2024)['LN'][0]:,.0f} personas.",
                              "Participación = SV / (SV + NV), sin las personas sin dato de voto."], tam=17)
    imagen(d, "participacion_nacional.png", 16.5, 4.2, 16)

    d = nueva(prs, "El método: celdas demográficas agrupadas por su comportamiento")
    texto(d, 1.5, 4, 30.8, 13, ["1. Unidad: celda = grupo de edad × sexo × tipo de sección × entidad.",
                                "2. Medida: participación de la celda menos la de su entidad, en cada elección.",
                                "3. Tres rasgos: nivel, brecha (presidencial menos intermedia) y tendencia.",
                                "4. K-means ponderado por lista nominal, semilla fija, k entre 4 y 6.",
                                "5. Validación: silhouette y estabilidad entre semillas y submuestras. Se eligió k = 5.",
                                "6. Presencia territorial: composición de cada sección, clave (entidad, sección)."], tam=19)

    d = nueva(prs, "Medir contra la propia entidad evita perfiles que son etiquetas de estado")
    tabla(d, [["Referencia", "Perfiles", "Entidades con más de 80 % en un perfil", "Varianza de la brecha explicada por la entidad"],
              ["Nacional (prueba de sensibilidad)", nac["k"], f"{nac['entidades_con_80pct_en_un_perfil']} de 32", f"{100 * nac['eta2_brecha']:.0f} %"],
              ["Entidad (análisis principal)", ent["k"], f"{ent['entidades_con_80pct_en_un_perfil']} de 32", f"{100 * ent['eta2_brecha']:.1f} %"]],
          1.5, 4.5, [9, 3.5, 9.5, 8.8])
    texto(d, 1.5, 9.5, 30.8, 7, ["Con la referencia nacional, algunas entidades quedaban casi enteras en un perfil.",
                                 "Una explicación posible es la concurrencia de elecciones locales; no se pudo verificar con estos datos."], tam=18)

    d = nueva(prs, "Cinco perfiles; los dos más jóvenes son la prioridad")
    imagen(d, "pesos.png", 1.5, 4, 17)
    texto(d, 19.5, 4.2, 13, 12, [f"{m[1]['pct_lista']:.0f} % de la lista tiene el perfil 1, el más numeroso.",
                                 f"El perfil 2 pesa {m[2]['pct_lista']:.0f} % y es el que más cayó en presidenciales.",
                                 "Los perfiles 3, 4 y 5 son de mayor edad y participan más."], tam=17)

    d = nueva(prs, "Se distinguen por cuánto participan y cómo cambia entre tipos de elección")
    imagen(d, "trayectorias.png", 1.5, 3.8, 19.5)
    texto(d, 22, 4.2, 10.5, 13, [f"Perfil 1: {m[1]['brecha_pp']:.0f} puntos más en presidenciales que en intermedias.",
                                 f"Perfil 2: de {100 * m[2]['part'][2012]:.0f} % (2012) a {100 * m[2]['part'][2024]:.0f} % (2024).",
                                 f"Perfil 4: de {100 * m[4]['part'][2009]:.0f} % (2009) a {100 * m[4]['part'][2024]:.0f} % (2024)."], tam=16)

    for titulo, perfiles in [("Prioridad alta: jóvenes con participación desigual o en caída", (1, 2)),
                             ("Prioridad media y mantener: adultez y vejez con participación alta", (3, 4, 5))]:
        d = nueva(prs, titulo)
        ancho = 30.8 / len(perfiles)
        for k, i in enumerate(perfiles):
            x = m[i]
            texto(d, 1.5 + k * ancho, 4, ancho - 0.8, 13,
                  [f"{i}. {x['nombre']}", f"{x['pct_lista']:.0f} % de la lista; mediana {x['mediana_edad']} años",
                   f"{x['pct_mujeres']:.0f} % mujeres; {x['pct_urbana']:.0f} % urbana",
                   f"Participación 2024: {100 * x['part'][2024]:.0f} %"], tam=16 if len(perfiles) == 3 else 18)
            d.shapes[-1].text_frame.paragraphs[0].runs[0].font.bold = True

    d = nueva(prs, "Una estrategia distinta para cada perfil, con indicador medible")
    tabla(d, [["Perfil", "Prioridad", "Objetivo", "Indicador"]] +
          [[f"{i}. {m[i]['nombre']}", contenido.ESTRATEGIAS[i]["prioridad"], contenido.ESTRATEGIAS[i]["objetivo"],
            contenido.ESTRATEGIAS[i]["indicador"]] for i in m],
          1.0, 3.8, [6.2, 2.4, 11.5, 11.7], tam=11)

    d = nueva(prs, "Limitaciones y próximos pasos")
    texto(d, 1.5, 4, 30.8, 13, ["Sin escolaridad en la fuente; se usa sexo y no género.",
                                "Los perfiles describen franjas de edad, no a las mismas personas en el tiempo.",
                                "Las causas (elecciones locales, movilidad laboral) son hipótesis por verificar.",
                                "Pendiente: contrastar la participación nacional con la plataforma CCPC.",
                                "Siguiente paso: piloto en dos o tres entidades con las estrategias 1 y 2."], tam=19)
    assert len(prs.slides) <= 10
    prs.save(ruta)


if __name__ == "__main__":
    destino = TABLAS.parent / "presentacion_ejecutiva.pptx"
    construir(destino)
    print("diapositivas:", len(Presentation(destino).slides))
