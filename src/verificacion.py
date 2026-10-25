"""Comprobaciones de los entregables RE-01 a RE-03 y de las reglas de antislop sobre su texto."""
from pathlib import Path

import pymupdf
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from configuracion import TABLAS

ENTREGA = TABLAS.parent
RAYA_LARGA = "—"


def texto_pdf(ruta: Path) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(ruta).pages)


def texto_pptx(ruta: Path) -> str:
    partes = []
    for diap in Presentation(ruta).slides:
        for forma in diap.shapes:
            if forma.has_text_frame:
                partes.append(forma.text_frame.text)
            if getattr(forma, "has_table", False) and forma.has_table:
                partes += [c.text for fila in forma.table.rows for c in fila.cells]
    return "\n".join(partes)


def fuentes_pdf(ruta: Path) -> set[str]:
    """Fuentes con las que se dibuja texto. Reportlab declara Helvetica como recurso aunque no la use."""
    usadas = set()
    for pagina in pymupdf.open(ruta):
        for bloque in pagina.get_text("dict")["blocks"]:
            for linea in bloque.get("lines", []):
                usadas |= {s["font"] for s in linea["spans"] if s["text"].strip()}
    return usadas


def resumen() -> dict:
    doc, deck = ENTREGA / "documento_tecnico.pdf", ENTREGA / "presentacion_ejecutiva.pptx"
    fichas = sorted((ENTREGA / "fichas").glob("*.png"))
    return {
        "paginas": len(PdfReader(doc).pages),
        "fuentes": fuentes_pdf(doc),
        "diapositivas": len(Presentation(deck).slides),
        "fichas": len(fichas),
        "dpi_fichas": [round(Image.open(f).info.get("dpi", (0, 0))[0]) for f in fichas],
        "rayas_largas": texto_pdf(doc).count(RAYA_LARGA) + texto_pptx(deck).count(RAYA_LARGA),
    }


if __name__ == "__main__":
    for clave, valor in resumen().items():
        print(clave, valor)
